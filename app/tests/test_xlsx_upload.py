"""Excel uploads: a workbook imports exactly as its CSV twin would — and a hostile
one is refused before `openpyxl` opens it.

Two promises. First, equivalence: the blueprint and results uploads accept
`.xlsx`, and a workbook's cells arrive as the strings Excel's own CSV export
would have written (`7` not `7.0`, `B` as `B`, an empty cell as a blank), so
every rule `results.py` applies to a CSV applies unchanged. Second, the attack
surface: a workbook is a zip from a stranger, so the member count, the declared
uncompressed size and any DTD are checked with `zipfile` first, and a forged
sheet size cannot make the reader pad rows to sixteen thousand cells.

Every workbook here is built in the test with openpyxl itself.
"""

from __future__ import annotations

import datetime as dt
import io
import re
import zipfile

import openpyxl
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile

from apps.events.models import Attempt
from apps.ingestion.models import QuestionTopicMap
from apps.ingestion.models import TestPaper as Paper
from apps.ingestion.services import tabular
from apps.ingestion.services.tabular import UnreadableFile, read_table
from tests import factories as f

pytestmark = pytest.mark.usefixtures("fast_password_hasher")

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def workbook(*rows, sheets: int = 1) -> bytes:
    """An .xlsx with `rows` on the first sheet. `None` is an empty cell."""
    book = openpyxl.Workbook()
    sheet = book.active
    for row in rows:
        sheet.append(list(row))
    for n in range(1, sheets):
        book.create_sheet(f"Other {n}").append(["not", "read"])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def upload(raw: bytes, name: str = "results.xlsx"):
    return SimpleUploadedFile(name, raw, content_type=XLSX)


# ----------------------------------------------------------- the conversion


def test_cells_arrive_as_the_strings_a_csv_would_carry():
    raw = workbook(
        ["Roll No", "Name", "Q1", "Q2", "Q3", "Q4", None],    # trailing empty header cell
        ["A-1", "Aarav", "B", "A", "D", 7],                    # int
        [None, None, None, None, None, None],                  # blank row: skipped
        ["A-2", "Bhavna", "c", None, "NR", 7.0],               # float that is whole
        [1041, "Chirag", 2, "b", "", 2.5],                     # a numeric roll number
    )
    table = read_table(raw)
    assert table.header == ["Roll No", "Name", "Q1", "Q2", "Q3", "Q4"]
    assert table.rows == [
        (2, ["A-1", "Aarav", "B", "A", "D", "7"]),
        # Line numbers are Excel's own row numbers, so row 3 (blank) is skipped
        # and an error on Bhavna names row 4, the row the sender sees.
        (4, ["A-2", "Bhavna", "c", "", "NR", "7"]),
        (5, ["1041", "Chirag", "2", "b", "", "2.5"]),
    ]


@pytest.mark.parametrize("value, text", [
    (None, ""), (7, "7"), (7.0, "7"), (2.5, "2.5"),
    # A long number Excel stored in E-notation comes back as a float; a
    # 16-digit enrolment number must not turn into "1.2345678901234e+15".
    (1.2345678901234e15, "1234567890123400"),
    (0.1 + 0.2, "0.3"),                                  # Excel's 15 digits, not float noise
    (True, "TRUE"), ("  B ", "B"),
    (dt.datetime(2026, 10, 4), "2026-10-04"), (dt.datetime(2026, 10, 4, 9, 30), "2026-10-04 09:30:00"),
])
def test_one_cell_as_excels_own_csv_export_would_write_it(value, text):
    assert tabular._cell_text(value) == text


def test_only_the_first_sheet_is_read():
    table = read_table(workbook(["roll_no", "Q1"], ["A-1", "B"], sheets=3))
    assert table.rows == [(2, ["A-1", "B"])]


def test_a_forged_sheet_size_does_not_pad_every_row(monkeypatch):
    """A `<dimension ref="A1:XFD1048576">` would have openpyxl pad each row to
    16,384 cells and add a million empty rows. The reader ignores it."""
    raw = workbook(["roll_no", "Q1"], ["A-1", "B"])
    forged = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as src, zipfile.ZipFile(forged, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data, n = re.subn(rb'<dimension ref="[^"]+"', b'<dimension ref="A1:XFD1048576"', data)
                assert n == 1
            dst.writestr(item, data)
    table = read_table(forged.getvalue())
    assert table.header == ["roll_no", "Q1"] and table.rows == [(2, ["A-1", "B"])]


# ------------------------------------------------------- through the upload


BLUEPRINT_ROWS = [
    ["question_id", "chapter", "correct_option", "answer"],
    ["Q1", "Kinematics", "B", None],
    ["Q2", "Kinematics", "A", None],
    ["Q3", "Thermodynamics", "D", None],
    ["Q4", "Limits", None, 7],                   # a numeric key, as a number
]


@pytest.fixture
def lab(cohort, mentor_client):
    students = {roll: f.make_student(cohort, name, roll_no=roll)
                for name, roll in [("Aarav", "A-1"), ("Bhavna", "A-2"), ("Chirag", "1041")]}
    response = mentor_client.post(
        "/api/papers/", {"name": "Mock 15", "exam": "JEE_MAIN", "held_on": "2026-10-04"},
        format="json",
    )
    assert response.status_code == 201, response.content
    paper = Paper.objects.get(pk=response.json()["id"])
    return type("Lab", (), dict(cohort=cohort, client=mentor_client, paper=paper, students=students))


def post_blueprint(lab, raw):
    return lab.client.post(f"/api/papers/{lab.paper.id}/blueprint/",
                           {"file": upload(raw, "blueprint.xlsx")}, format="multipart")


def post_results(lab, raw, **extra):
    return lab.client.post("/api/ingest/results/",
                           {"file": upload(raw), "paper": lab.paper.id, **extra}, format="multipart")


def test_a_blueprint_workbook_maps_every_question(lab):
    response = post_blueprint(lab, workbook(*BLUEPRINT_ROWS))
    assert response.status_code == 200, response.content
    assert response.json()["created"] == 4
    q4 = QuestionTopicMap.objects.get(test_paper=lab.paper, question_id="Q4")
    assert q4.answer_type == "numerical" and q4.numeric_answer == 7.0
    q1 = QuestionTopicMap.objects.get(test_paper=lab.paper, question_id="Q1")
    assert q1.options.get(is_correct=True).label == "B"


def test_a_wide_results_workbook_is_marked_like_the_csv(lab):
    assert post_blueprint(lab, workbook(*BLUEPRINT_ROWS)).status_code == 200
    raw = workbook(
        ["Roll No", "Name", "Q1", "Q2", "Q3", "Q4", "Total"],
        ["A-1", "Aarav", "B", "A", "D", 7, 16],
        ["A-2", "Bhavna", "C", "-", "NR", 8, -2],
        [1041, "Chirag", 2, "b", None, 7.0, 7],       # roll typed as a number; a blank
    )
    response = post_results(lab, raw)
    assert response.status_code == 201, response.content
    body = response.json()
    assert (body["layout"], body["students_ingested"], body["attempts_created"]) == ("wide", 3, 12)
    assert body["ignored_columns"] == ["Name", "Total"]

    def statuses(roll):
        rows = Attempt.objects.filter(student=lab.students[roll], test_paper=lab.paper)
        return {a.question_id: (a.status, a.marks, a.chosen_option) for a in rows}

    assert statuses("A-1") == {"Q1": ("correct", 4, "B"), "Q2": ("correct", 4, "A"),
                               "Q3": ("correct", 4, "D"), "Q4": ("correct", 4, "")}
    assert statuses("A-2") == {"Q1": ("wrong", -1, "C"), "Q2": ("blank", 0, ""),
                               "Q3": ("not_reached", 0, ""), "Q4": ("wrong", -1, "")}
    # The digit 2 is option B; the empty cell is a blank, never a guess.
    assert statuses("1041") == {"Q1": ("correct", 4, "B"), "Q2": ("wrong", -1, "B"),
                                "Q3": ("blank", 0, ""), "Q4": ("correct", 4, "")}


def test_a_long_results_workbook_carries_time_per_question(lab):
    assert post_blueprint(lab, workbook(*BLUEPRINT_ROWS)).status_code == 200
    raw = workbook(
        ["roll_no", "question_id", "answer", "time_spent"],
        ["A-1", "Q1", "B", 95],
        ["A-1", "Q2", "C", 120.0],
        ["A-1", "Q4", 7, None],
    )
    response = post_results(lab, raw)
    assert response.status_code == 201, response.content
    assert response.json()["layout"] == "long"
    got = {a.question_id: (a.status, a.time_spent)
           for a in Attempt.objects.filter(student=lab.students["A-1"], test_paper=lab.paper)}
    assert got == {"Q1": ("correct", 95), "Q2": ("wrong", 120), "Q4": ("correct", None)}


def test_a_bad_cell_in_a_workbook_names_the_excel_row(lab):
    assert post_blueprint(lab, workbook(*BLUEPRINT_ROWS)).status_code == 200
    raw = workbook(["roll_no", "Q1"], ["A-1", "B"], [None], ["A-2", "E"])
    body = post_results(lab, raw).json()
    assert body["students_failed"] == 1
    assert body["errors"][0]["row"] == 4 and "'E' is not an option" in body["errors"][0]["message"]


# ---------------------------------------------------------- the refusals


def _zip(members: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in members.items():
            z.writestr(name, data)
    return buffer.getvalue()


SKELETON = {
    "[Content_Types].xml": b'<?xml version="1.0"?><Types/>',
    "xl/workbook.xml": b'<?xml version="1.0"?><workbook/>',
}


@pytest.fixture
def openpyxl_must_not_open(monkeypatch):
    """The guard's whole point is that it runs first."""
    def refuse(*args, **kwargs):
        raise AssertionError("openpyxl was handed a file the zip guard should have refused")
    monkeypatch.setattr(openpyxl, "load_workbook", refuse)


def test_a_zip_bomb_is_refused_before_openpyxl_sees_it(openpyxl_must_not_open):
    # 26 MB of one byte deflates to a few kilobytes: well under the 2 MB upload
    # limit, well over what a results sheet unpacks to.
    raw = _zip({**SKELETON, "xl/worksheets/sheet1.xml": b" " * 26_000_000})
    assert len(raw) < tabular.MAX_BYTES
    with pytest.raises(UnreadableFile, match="unpacks to 26 MB"):
        read_table(raw)


def test_an_archive_with_too_many_members_is_refused(openpyxl_must_not_open):
    raw = _zip({**SKELETON, **{f"xl/media/{i}.bin": b"x" for i in range(tabular.MAX_XLSX_MEMBERS)}})
    with pytest.raises(UnreadableFile, match="parts inside it"):
        read_table(raw)


def test_a_dtd_is_refused_before_any_xml_is_parsed(openpyxl_must_not_open):
    laughs = b'<?xml version="1.0"?><!DOCTYPE lolz [<!ENTITY lol "lol">]><worksheet>&lol;</worksheet>'
    with pytest.raises(UnreadableFile, match="markup Excel never writes"):
        read_table(_zip({**SKELETON, "xl/worksheets/sheet1.xml": laughs}))


def test_a_zip_that_is_not_a_workbook_is_refused(openpyxl_must_not_open):
    with pytest.raises(UnreadableFile, match="not an Excel workbook"):
        read_table(_zip({"notes.txt": b"hello"}))


def test_a_row_reaching_far_to_the_right_is_refused():
    book = openpyxl.Workbook()
    book.active.append(["roll_no", "Q1"])
    book.active.cell(row=2, column=tabular.MAX_COLUMNS + 5, value="x")
    buffer = io.BytesIO()
    book.save(buffer)
    with pytest.raises(UnreadableFile, match="reaches past column"):
        read_table(buffer.getvalue())


def test_reading_stops_at_the_row_limit(monkeypatch):
    monkeypatch.setattr(tabular, "MAX_ROWS", 3)
    raw = workbook(["roll_no", "Q1"], *[[f"A-{i}", "B"] for i in range(5)])
    with pytest.raises(UnreadableFile, match="more than 3 rows"):
        read_table(raw)


def test_an_empty_workbook_is_a_sentence():
    with pytest.raises(UnreadableFile, match="is empty"):
        read_table(workbook())


def test_legacy_xls_is_still_refused_with_what_to_do(lab):
    response = lab.client.post(
        f"/api/papers/{lab.paper.id}/blueprint/",
        {"file": SimpleUploadedFile("old.xls", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 64)},
        format="multipart",
    )
    assert response.status_code == 400
    detail = response.json()["detail"]
    assert ".xls" in detail and "Save As" in detail and ".xlsx" in detail


def test_a_refused_workbook_is_a_400_with_a_sentence_on_both_uploads(lab):
    bomb = _zip({**SKELETON, "xl/worksheets/sheet1.xml": b" " * 26_000_000})
    blueprint = post_blueprint(lab, bomb)
    results = post_results(lab, bomb)
    for response in (blueprint, results):
        assert response.status_code == 400 and "unpacks to" in response.json()["detail"]
